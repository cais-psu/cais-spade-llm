"""Write the case-study material flow as an editable PowerPoint block diagram.

Each resource is one rounded block; part-flow arrows are PowerPoint connectors
glued to the blocks, so they follow a block when it is moved. Slide 1 shows the
flow; slide 2 adds the four failure scenarios. Slide 3 labels the Gazebo view
(``gazebo_view.png``) with the same category colours as the blocks.

The flow follows the recovery-framework configuration: KMR delivers the KET
pegs to M1 and the RGOCG pegs to M2, the machine robots load the Conveyor,
ur5e-3 takes pegs from the Buffer into assembly and moves the finished board
to Exit, and ur5e-4 assembles the printed gears.

Text is sized for a 13.333 in slide that is scaled to a double-column journal
figure (about 0.54x): 16 pt names and 12 pt details print at about 8.6 pt
and 6.4 pt.

Requires ``python-pptx`` (not a project dependency):

    python case_study_flow_pptx.py
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

OUTPUT_PATH = Path(__file__).resolve().parent / "case_study_flow.pptx"
GAZEBO_VIEW_PATH = Path(__file__).resolve().parent / "gazebo_view.png"
GAZEBO_VIEW_SIZE_PX = (1422, 583)
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 6.1
FONT = "Arial"
SYMBOL_FONT = "Segoe UI Symbol"
NAME_PT = 16
DETAIL_PT = 12
LEGEND_PT = 12

BLOCK_H = 1.06
COLUMN_GAP = 0.40
ROW_A, ROW_C, ROW_B, ROW_E = 0.83, 2.31, 3.79, 5.27

INK = "2B2F36"
MUTED = "5D6673"
FLOW = "1F5FA8"
FAILURE = "D62728"
WHITE = "FFFFFF"

# (fill, border) for each resource category.
CATEGORY_STYLE = {
    "storage": ("DCE8F5", "3D5A80"),
    "transport": ("FBE3C2", "B8661D"),
    "process": ("E3E6EB", "4A4F57"),
    "robot": ("D7EEE7", "2A8C7A"),
    "product": ("FFFFFF", "2B2F36"),
}
CATEGORY_LABELS = (
    ("storage", "Storage / buffer"),
    ("transport", "Transport"),
    ("process", "Processing"),
    ("robot", "Robot arm"),
    ("product", "Product"),
)

# Part symbol and colour, shown in each block for the parts that resource handles.
PART_SYMBOLS = {
    "KET": ("■", "3D6FB6"),
    "RGOCG": ("●", "8B5CB8"),
    "gear": ("✹", "E0843C"),
    "board": ("▣", "2B2F36"),
}
PART_LABELS = (
    ("KET", "KET square peg"),
    ("RGOCG", "RGOCG round peg"),
    ("gear", "Printed gear"),
    ("board", "Finished board"),
)
PEGS = (("KET", 4), ("RGOCG", 4))

FAILURE_SCENARIOS = (
    "Conveyor breakdown",
    "ur5e-1 breakdown",
    "Machining breakdown during part processing",
    "Part slippage",
)

# Connection-site indices of a rectangle preset.
TOP, LEFT, BOTTOM, RIGHT = 0, 1, 2, 3


@dataclass(frozen=True)
class Block:
    """One resource block of the flow diagram."""

    name: str
    detail: str
    category: str
    column: int
    row: float
    width: float
    parts: tuple[tuple[str, int], ...]


# Part responsibilities follow recovery_framework_gazebo.json: M1/ur5e-1 handle the KET
# pegs, M2/ur5e-2 the RGOCG pegs, ur5e-3 the Buffer and Exit, and ur5e-4 the printed gears.
BLOCKS = (
    Block("Storage", "raw pegs", "storage", 0, ROW_A, 1.42, PEGS),
    Block("KUKA KMR", "mobile robot", "transport", 0, ROW_C, 1.42, PEGS),
    Block("M1", "CNC trimming", "process", 1, ROW_A, 1.50, (("KET", 4),)),
    Block("M2", "CNC trimming", "process", 1, ROW_B, 1.50, (("RGOCG", 4),)),
    Block("UR5e-1", "machine tending", "robot", 2, ROW_A, 1.60, (("KET", 4),)),
    Block("UR5e-2", "machine tending", "robot", 2, ROW_B, 1.60, (("RGOCG", 4),)),
    Block("Conveyor", "belt transport", "transport", 3, ROW_C, 1.25, PEGS),
    Block("Buffer", "4 zones", "storage", 4, ROW_C, 1.25, PEGS),
    Block("Printer", "3D printing", "process", 4, ROW_E, 1.25, (("gear", 3),)),
    Block("UR5e-3", "assembly", "robot", 5, ROW_C, 1.60, PEGS + (("board", 1),)),
    Block("Assembly Board", "assembly_board-v1", "product", 5, ROW_B, 1.97, PEGS + (("gear", 3),)),
    Block("UR5e-4", "gear assembly", "robot", 5, ROW_E, 1.60, (("gear", 3),)),
    Block("Exit", "finished product", "storage", 6, ROW_C, 1.41, (("board", 1),)),
)
BLOCKS_BY_NAME = {block.name: block for block in BLOCKS}

# Label boxes on the Gazebo view as (name, left, right, top, bottom) in screenshot pixels.
PHOTO_LABELS = (
    ("KUKA KMR", 128, 367, 22, 82),
    ("M1", 421, 506, 16, 77),
    ("M2", 714, 799, 16, 77),
    ("UR5e-1", 486, 638, 187, 247),
    ("UR5e-2", 799, 950, 156, 216),
    ("UR5e-3", 1066, 1217, 156, 217),
    ("Buffer", 914, 1055, 221, 280),
    ("Exit", 1206, 1315, 245, 306),
    ("Storage", 14, 211, 279, 340),
    ("Conveyor", 519, 715, 343, 404),
    ("Assembly Board", 883, 1065, 345, 451),
    ("Printer", 1249, 1401, 377, 438),
    ("UR5e-4", 1055, 1206, 455, 515),
)

# (source, source site, target, target site, connector type)
FLOWS = (
    ("Storage", BOTTOM, "KUKA KMR", TOP, "straight"),
    ("KUKA KMR", RIGHT, "M1", LEFT, "elbow"),
    ("KUKA KMR", RIGHT, "M2", LEFT, "elbow"),
    ("M1", RIGHT, "UR5e-1", LEFT, "straight"),
    ("M2", RIGHT, "UR5e-2", LEFT, "straight"),
    ("UR5e-1", RIGHT, "Conveyor", LEFT, "elbow"),
    ("UR5e-2", RIGHT, "Conveyor", LEFT, "elbow"),
    ("Conveyor", RIGHT, "Buffer", LEFT, "straight"),
    ("Buffer", RIGHT, "UR5e-3", LEFT, "straight"),
    ("UR5e-3", BOTTOM, "Assembly Board", TOP, "straight"),
    ("Printer", RIGHT, "UR5e-4", LEFT, "straight"),
    ("UR5e-4", TOP, "Assembly Board", BOTTOM, "straight"),
    ("Assembly Board", RIGHT, "Exit", BOTTOM, "corner"),
)


def _column_centers() -> list[float]:
    column_count = max(block.column for block in BLOCKS) + 1
    widths = [max(block.width for block in BLOCKS if block.column == column) for column in range(column_count)]
    left = (SLIDE_WIDTH_IN - sum(widths) - COLUMN_GAP * (column_count - 1)) / 2
    centers = []
    for width in widths:
        centers.append(left + width / 2)
        left += width + COLUMN_GAP
    return centers


COLUMN_CENTERS = _column_centers()


def _center(name: str) -> tuple[float, float]:
    block = BLOCKS_BY_NAME[name]
    return COLUMN_CENTERS[block.column], block.row


def _left_edge(name: str) -> float:
    return _center(name)[0] - BLOCKS_BY_NAME[name].width / 2


def _style_run(run, size_pt: float, *, bold: bool = False, color: str = INK, font: str = FONT) -> None:
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_symbol_run(paragraph, part: str) -> None:
    symbol, color = PART_SYMBOLS[part]
    run = paragraph.add_run()
    run.text = symbol
    _style_run(run, DETAIL_PT + 2, color=color, font=SYMBOL_FONT)


def _add_part_runs(paragraph, parts: tuple[tuple[str, int], ...]) -> None:
    for index, (part, count) in enumerate(parts):
        if index:
            spacer = paragraph.add_run()
            spacer.text = "   "
            _style_run(spacer, DETAIL_PT)
        _add_symbol_run(paragraph, part)
        run = paragraph.add_run()
        run.text = f" ×{count}"
        _style_run(run, DETAIL_PT, bold=True)


def _clear_margins(frame) -> None:
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0


def _style_box(shape, category: str) -> None:
    fill, border = CATEGORY_STYLE[category]
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(border)
    shape.line.width = Pt(2.25 if category == "product" else 1.5)


def _add_block(shapes, block: Block):
    x, y = _center(block.name)
    shape = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - block.width / 2), Inches(y - BLOCK_H / 2),
                             Inches(block.width), Inches(BLOCK_H))
    shape.name = block.name
    shape.adjustments[0] = 0.13
    _style_box(shape, block.category)

    frame = shape.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, (text, size_pt, bold, color) in enumerate(((block.name, NAME_PT, True, INK),
                                                          (block.detail, DETAIL_PT, False, MUTED))):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text
        _style_run(run, size_pt, bold=bold, color=color)
    parts_paragraph = frame.add_paragraph()
    parts_paragraph.alignment = PP_ALIGN.CENTER
    parts_paragraph.space_before = Pt(3)
    _add_part_runs(parts_paragraph, block.parts)
    return shape


def _add_arrowhead(connector) -> None:
    connector.line.color.rgb = RGBColor.from_string(FLOW)
    connector.line.width = Pt(2.0)
    tail = etree.SubElement(connector.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _add_flow(shapes, blocks: dict, source: str, source_site: int, target: str, target_site: int,
              kind: str) -> None:
    connector_type = MSO_CONNECTOR.STRAIGHT if kind == "straight" else MSO_CONNECTOR.ELBOW
    connector = shapes.add_connector(connector_type, 0, 0, Inches(1), Inches(1))
    connector.begin_connect(blocks[source], source_site)
    connector.end_connect(blocks[target], target_site)
    connector.name = f"Flow {source} to {target}"
    if kind == "corner":
        # One bend: leave the board sideways, then rise into Exit from below.
        connector._element.spPr.find(qn("a:prstGeom")).set("prst", "bentConnector2")
    _add_arrowhead(connector)


def _add_label(shapes, x: float, y: float, text: str, name: str, *, align: str = "center",
               size_pt: float = LEGEND_PT, bold: bool = False, color: str = INK, width: float = 1.6,
               part: str | None = None):
    height = size_pt * 1.4 / 72
    left = {"left": x, "right": x - width}.get(align, x - width / 2)
    box = shapes.add_textbox(Inches(left), Inches(y - height / 2), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.CENTER)
    if part is not None:
        _add_symbol_run(paragraph, part)
        text = f"  {text}"
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size_pt, bold=bold, color=color)
    return box


def _add_badge(shapes, x: float, y: float, number: int, name: str, radius: float = 0.14) -> None:
    badge = shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - radius), Inches(y - radius), Inches(2 * radius),
                             Inches(2 * radius))
    badge.name = name
    badge.shadow.inherit = False
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor.from_string(FAILURE)
    badge.line.color.rgb = RGBColor.from_string(WHITE)
    badge.line.width = Pt(1.5)
    frame = badge.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    _style_run(run, 12, bold=True, color=WHITE)


def _add_exit_transfer_label(shapes) -> None:
    exit_x, exit_y = _center("Exit")
    _add_label(shapes, exit_x - 0.08, (ROW_B + exit_y + BLOCK_H / 2) / 2, "by UR5e-3", "Label Exit transfer",
               align="right", color=MUTED, width=1.0)


def _add_legend(shapes) -> None:
    legend = shapes.add_group_shape()
    legend.name = "Legend"
    left = _left_edge("Storage")
    column_offsets, row_step = (0.0, 2.0, 3.95, 5.55), 0.38
    for index, (part, text) in enumerate(PART_LABELS):
        _add_label(legend.shapes, left + column_offsets[index], ROW_E + row_step, text, f"Legend part {text}",
                   align="left", width=1.8, part=part)
    entries = list(CATEGORY_LABELS) + [("flow", "Part flow")]
    for index, (category, text) in enumerate(entries):
        x = left + column_offsets[index % 4]
        y = ROW_E - row_step + (index // 4) * row_step
        if category == "flow":
            connector = legend.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                                    Inches(x + 0.44), Inches(y))
            connector.name = "Legend part flow"
            _add_arrowhead(connector)
        else:
            swatch = legend.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y - 0.13),
                                             Inches(0.44), Inches(0.26))
            swatch.name = f"Legend {text}"
            swatch.adjustments[0] = 0.2
            _style_box(swatch, category)
        _add_label(legend.shapes, x + 0.58, y, text, f"Legend text {text}", align="left")
    legend._element.recalculate_extents()


def _add_failures(shapes) -> None:
    markers = shapes.add_group_shape()
    markers.name = "Failure markers"
    for number, block_name in ((1, "Conveyor"), (2, "UR5e-1"), (3, "M1")):
        x, y = _center(block_name)
        _add_badge(markers.shapes, x + BLOCKS_BY_NAME[block_name].width / 2 - 0.03, y - BLOCK_H / 2 + 0.03,
                   number, f"Failure {number} {FAILURE_SCENARIOS[number - 1]}")
    # Part slippage happens while ur5e-4 places a gear onto the board.
    board_x, _ = _center("Assembly Board")
    _add_badge(markers.shapes, board_x + 0.27, (ROW_B + ROW_E) / 2, 4, "Failure 4 Part slippage")
    markers._element.recalculate_extents()

    legend = shapes.add_group_shape()
    legend.name = "Failure scenarios legend"
    left = _left_edge("Conveyor")
    _add_label(legend.shapes, left, ROW_A - 0.36, "Failure scenarios", "Failure scenarios title", align="left",
               size_pt=13, bold=True, width=2.2)
    for number, scenario in enumerate(FAILURE_SCENARIOS, start=1):
        x = left + ((number - 1) % 2) * 4.0
        y = ROW_A + 0.02 + ((number - 1) // 2) * 0.36
        _add_badge(legend.shapes, x + 0.13, y, number, f"Failure legend badge {number}", radius=0.12)
        _add_label(legend.shapes, x + 0.36, y, scenario, f"Failure legend text {number}", align="left", width=3.5)
    legend._element.recalculate_extents()


def _draw_slide(shapes, show_failures: bool) -> None:
    blocks = {block.name: _add_block(shapes, block) for block in BLOCKS}
    for flow in FLOWS:
        _add_flow(shapes, blocks, *flow)
    _add_exit_transfer_label(shapes)
    _add_legend(shapes)
    if show_failures:
        _add_failures(shapes)


def _draw_photo_slide(shapes) -> None:
    image_w_px, image_h_px = GAZEBO_VIEW_SIZE_PX
    inches_per_px = SLIDE_WIDTH_IN / image_w_px
    top = (SLIDE_HEIGHT_IN - image_h_px * inches_per_px) / 2
    picture = shapes.add_picture(str(GAZEBO_VIEW_PATH), 0, Inches(top), width=Inches(SLIDE_WIDTH_IN))
    picture.name = "Gazebo view"

    for name, left_px, right_px, top_px, bottom_px in PHOTO_LABELS:
        label = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_px * inches_per_px),
                                 Inches(top + top_px * inches_per_px), Inches((right_px - left_px) * inches_per_px),
                                 Inches((bottom_px - top_px) * inches_per_px))
        label.name = f"Label {name}"
        label.adjustments[0] = 0.13
        _style_box(label, BLOCKS_BY_NAME[name].category)
        frame = label.text_frame
        _clear_margins(frame)
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for index, line in enumerate(name.split(" ") if name == "Assembly Board" else [name]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = line
            _style_run(run, NAME_PT, bold=True)


def build_presentation() -> Presentation:
    """Build the flow diagram deck.

    Returns:
        A presentation with the flow slide, the failure-scenario slide, and
        the labelled Gazebo view.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = presentation.slide_layouts[6]
    for show_failures in (False, True):
        _draw_slide(presentation.slides.add_slide(blank_layout).shapes, show_failures)
    _draw_photo_slide(presentation.slides.add_slide(blank_layout).shapes)
    return presentation


def main() -> None:
    """Write ``case_study_flow.pptx`` next to this script."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    build_presentation().save(OUTPUT_PATH)
    logger.info("Wrote %s", OUTPUT_PATH)


if __name__ == "__main__":
    main()
