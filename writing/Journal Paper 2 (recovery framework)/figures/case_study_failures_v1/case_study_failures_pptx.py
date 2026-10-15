"""Write the four failure scenarios as editable PowerPoint figures.

``case_study_failures.pptx`` shows one panel per failure: the Gazebo snapshot
next to a card with the failed resource, the injection checkpoint, the state
left at failure, and the disrupted capability. ``case_study_failure_timeline.pptx``
places the four injection checkpoints on the nominal part routes.

Card text follows the official failure captures (``observations.json``):
``capture_1b3f3d1d729a4701b94cff514c783130`` (Conveyor breakdown),
``capture_0a7fde8074b6426586b204204f2b3417`` (ur5e-1 breakdown),
``capture_8beafdb8a01e444894cfaf41d29921fd`` (Machining breakdown during part
processing), and ``cf647a0930c64f018d672a27904811b0`` (Part slippage).
Colours, numbering, and resource labels match ``case_study_flow_v1``.

Requires ``python-pptx`` (not a project dependency):

    python case_study_failures_pptx.py
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from matplotlib.image import imread, imsave
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

FIGURE_DIR = Path(__file__).resolve().parent
COMPOSITE_PATH = FIGURE_DIR / "failure_composite.png"
PANELS_PATH = FIGURE_DIR / "case_study_failures.pptx"
TIMELINE_PATH = FIGURE_DIR / "case_study_failure_timeline.pptx"
SLIDE_WIDTH_IN = 13.333
PANELS_HEIGHT_IN = 7.8
TIMELINE_HEIGHT_IN = 3.75
FONT = "Arial"
SYMBOL_FONT = "Segoe UI Symbol"

INK = "2B2F36"
MUTED = "5D6673"
FLOW = "1F5FA8"
FAILURE = "D62728"
WHITE = "FFFFFF"
PANEL_EDGE = "D5DAE1"
CARD_FILL = "F4F6F8"

# (fill, border) for each resource category, as in the flow block diagram.
CATEGORY_STYLE = {
    "storage": ("DCE8F5", "3D5A80"),
    "transport": ("FBE3C2", "B8661D"),
    "process": ("E3E6EB", "4A4F57"),
    "robot": ("D7EEE7", "2A8C7A"),
    "product": ("FFFFFF", "2B2F36"),
}
PART_SYMBOLS = {"KET": ("■", "3D6FB6"), "gear": ("✹", "E0843C")}

# Connection-site indices of a rectangle preset.
LEFT, RIGHT = 1, 3


@dataclass(frozen=True)
class Failure:
    """One injected failure scenario and its observed state."""

    number: int
    scenario: str
    resource: str
    category: str
    injected: str
    state: tuple[str, ...]
    disrupted: str
    crop_px: tuple[int, int, int, int]

    @property
    def image_path(self) -> Path:
        return FIGURE_DIR / f"failure_{self.number}.png"


# Numbered as in EXPERIMENT_DESIGN_AND_PLAN.md; crops are (left, top, right, bottom) in the composite.
FAILURES = (
    Failure(1, "Conveyor breakdown", "Conveyor", "transport",
            "UR5e-1 holds the trimmed peg, before release onto the Conveyor",
            ("KET4_Square_4mm in UR5e-1", "gear_small in UR5e-4"),
            "transport to Buffer", (14, 370, 461, 682)),
    Failure(2, "ur5e-1 breakdown", "UR5e-1", "robot",
            "M1 has finished; UR5e-1 has approached but not grasped",
            ("trimmed KET4_Square_4mm in M1", "UR5e-1 stopped in M1's front access, gripper empty"),
            "M1 unloading", (485, 4, 932, 315)),
    Failure(3, "Machining breakdown during part processing", "M1", "process",
            "18% into M1's 5 s trim job",
            ("KET4_Square_4mm partly trimmed in M1", "RGOCG4-50_Round_4mm still trimming in M2"),
            "square trimming at M1", (14, 4, 461, 315)),
    Failure(4, "Part slippage", "UR5e-4", "robot",
            "gear_small half lowered onto the Assembly Board",
            ("gear_small upright in UR5e-3's region", "KET4_Square_4mm held by UR5e-3"),
            "gear_small placement", (485, 370, 932, 682)),
)


def _crop_snapshots() -> None:
    composite = imread(COMPOSITE_PATH)
    for failure in FAILURES:
        left, top, right, bottom = failure.crop_px
        imsave(failure.image_path, composite[top:bottom, left:right])


def _clear_margins(frame) -> None:
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0


def _style_run(run, size_pt: float, *, bold: bool = False, color: str = INK, font: str = FONT) -> None:
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _box(shapes, left: float, top: float, width: float, height: float, name: str, *, fill: str | None,
         line: str | None, line_pt: float = 1.5, corner: float = 0.1):
    shape = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    shape.adjustments[0] = corner
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(line_pt)
    return shape


def _set_lines(shape, lines: list[tuple[str, float, bool, str]], *, align=PP_ALIGN.CENTER) -> None:
    frame = shape.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, (text, size_pt, bold, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        _style_run(run, size_pt, bold=bold, color=color)


def _resource_box(shapes, left: float, top: float, width: float, height: float, name: str, category: str,
                  lines: list[tuple[str, float, bool, str]], *, failed: bool = False):
    fill, border = CATEGORY_STYLE[category]
    shape = _box(shapes, left, top, width, height, name, fill=fill,
                 line=FAILURE if failed else border, line_pt=2.5 if failed else 1.5, corner=0.16)
    _set_lines(shape, lines)
    return shape


def _text(shapes, left: float, top: float, width: float, height: float, name: str, text: str, *,
          size_pt: float, bold: bool = False, color: str = INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    box = shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size_pt, bold=bold, color=color)
    return box


def _badge(shapes, x: float, y: float, number: int, name: str, radius: float = 0.16, size_pt: float = 13) -> None:
    badge = shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - radius), Inches(y - radius), Inches(2 * radius),
                             Inches(2 * radius))
    badge.name = name
    badge.shadow.inherit = False
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor.from_string(FAILURE)
    badge.line.color.rgb = RGBColor.from_string(WHITE)
    badge.line.width = Pt(1.5)
    _set_lines(badge, [(str(number), size_pt, True, WHITE)])


def _arrow(shapes, source, target, name: str, connector_type=MSO_CONNECTOR.STRAIGHT) -> None:
    connector = shapes.add_connector(connector_type, 0, 0, Inches(1), Inches(1))
    connector.begin_connect(source, RIGHT)
    connector.end_connect(target, LEFT)
    connector.name = name
    connector.line.color.rgb = RGBColor.from_string(FLOW)
    connector.line.width = Pt(2.0)
    tail = etree.SubElement(connector.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _red_line(shapes, x: float, top: float, bottom: float, name: str) -> None:
    line = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(top), Inches(x), Inches(bottom))
    line.name = name
    line.line.color.rgb = RGBColor.from_string(FAILURE)
    line.line.width = Pt(2.25)


# ---------------------------------------------------------------------------------------------
# Figure 1: one panel per failure
# ---------------------------------------------------------------------------------------------

MARGIN = 0.3
GUTTER = 0.3
CELL_W = (SLIDE_WIDTH_IN - 2 * MARGIN - GUTTER) / 2
CELL_H = (PANELS_HEIGHT_IN - 2 * MARGIN - GUTTER) / 2
IMAGE_H = 2.3
CARD_PT = 12


def _add_card_text(shapes, left: float, top: float, width: float, height: float, failure: Failure) -> None:
    box = shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = f"Failure {failure.number} details"
    frame = box.text_frame
    _clear_margins(frame)
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    # Each observed fact starts its own line, so part identifiers are not split across lines.
    rows = (("Injected", (failure.injected,)), ("At failure", failure.state), ("Disrupted", (failure.disrupted,)))
    for index, (label, facts) in enumerate(rows):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(5)
        label_run = paragraph.add_run()
        label_run.text = f"{label}  "
        _style_run(label_run, CARD_PT, bold=True)
        for fact_index, fact in enumerate(facts):
            if fact_index:
                paragraph.add_line_break()
            value_run = paragraph.add_run()
            value_run.text = fact
            _style_run(value_run, CARD_PT, color=MUTED)


def _add_failure_panel(shapes, failure: Failure, left: float, top: float) -> None:
    panel = shapes.add_group_shape()
    panel.name = f"Failure {failure.number} {failure.scenario}"
    group = panel.shapes
    _box(group, left, top, CELL_W, CELL_H, "Panel frame", fill=WHITE, line=PANEL_EDGE, line_pt=1.0, corner=0.04)
    _badge(group, left + 0.32, top + 0.3, failure.number, "Panel badge")
    _text(group, left + 0.58, top + 0.1, CELL_W - 0.7, 0.4, "Panel title", failure.scenario, size_pt=16, bold=True)

    image_top = top + 0.58
    image_w = IMAGE_H * (failure.crop_px[2] - failure.crop_px[0]) / (failure.crop_px[3] - failure.crop_px[1])
    picture = group.add_picture(str(failure.image_path), Inches(left + 0.15), Inches(image_top),
                                height=Inches(IMAGE_H))
    picture.name = "Gazebo snapshot"

    card_left = left + 0.15 + image_w + 0.18
    card_w = left + CELL_W - 0.15 - card_left
    card_h = top + CELL_H - 0.15 - image_top
    _box(group, card_left, image_top, card_w, card_h, "Card", fill=CARD_FILL, line=None, corner=0.05)
    _text(group, card_left + 0.14, image_top + 0.12, 0.8, 0.36, "Failed label", "Failed", size_pt=CARD_PT, bold=True)
    chip_w = 0.32 + 0.12 * len(failure.resource)
    _resource_box(group, card_left + 0.86, image_top + 0.1, chip_w, 0.4, "Failed resource", failure.category,
                  [(failure.resource, 14, True, INK)], failed=True)
    _add_card_text(group, card_left + 0.14, image_top + 0.66, card_w - 0.26, card_h - 0.74, failure)
    panel._element.recalculate_extents()


def build_panels() -> Presentation:
    """Build the 2x2 failure panel figure.

    Returns:
        A one-slide presentation with one panel per failure scenario.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(PANELS_HEIGHT_IN)
    shapes = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes
    for index, failure in enumerate(FAILURES):
        left = MARGIN + (index % 2) * (CELL_W + GUTTER)
        top = MARGIN + (index // 2) * (CELL_H + GUTTER)
        _add_failure_panel(shapes, failure, left, top)
    return presentation


# ---------------------------------------------------------------------------------------------
# Figure 2: injection checkpoints on the nominal part routes
# ---------------------------------------------------------------------------------------------

STEP_W = 1.22
STEP_H = 0.66
STEP_GAP = 0.27
LANE_LABEL_W = 2.35
KET_LANE_Y = 1.18
GEAR_LANE_Y = 2.62

# (column, resource, action, category) for each nominal step of the two failure-run parts.
KET_STEPS = (
    (0, "KUKA KMR", "deliver to M1", "transport"),
    (1, "M1", "trim (5 s)", "process"),
    (2, "UR5e-1", "pick", "robot"),
    (3, "UR5e-1", "place", "robot"),
    (4, "Conveyor", "to Buffer", "transport"),
    (5, "UR5e-3", "assemble", "robot"),
)
GEAR_STEPS = (
    (3, "Printer", "gear ready", "process"),
    (4, "UR5e-4", "pick", "robot"),
    (5, "UR5e-4", "place", "robot"),
)


def _step_left(column: int) -> float:
    return MARGIN + LANE_LABEL_W + column * (STEP_W + STEP_GAP)


def _add_lane_label(shapes, y: float, part: str, name: str, note: str) -> None:
    box = shapes.add_textbox(Inches(MARGIN), Inches(y - 0.32), Inches(LANE_LABEL_W - 0.1), Inches(0.64))
    box.name = f"Lane {name}"
    frame = box.text_frame
    _clear_margins(frame)
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    symbol, color = PART_SYMBOLS[part]
    first = frame.paragraphs[0]
    symbol_run = first.add_run()
    symbol_run.text = f"{symbol} "
    _style_run(symbol_run, 15, color=color, font=SYMBOL_FONT)
    name_run = first.add_run()
    name_run.text = name
    _style_run(name_run, 14, bold=True)
    second = frame.add_paragraph()
    note_run = second.add_run()
    note_run.text = note
    _style_run(note_run, 12, color=MUTED)


def _add_lane(shapes, y: float, steps: tuple) -> list:
    boxes = []
    for column, resource, action, category in steps:
        boxes.append(_resource_box(shapes, _step_left(column), y - STEP_H / 2, STEP_W, STEP_H,
                                   f"{resource} {action}", category,
                                   [(resource, 14, True, INK), (action, 12, False, MUTED)]))
    for source, target in zip(boxes, boxes[1:]):
        _arrow(shapes, source, target, f"Flow {source.name} to {target.name}")
    return boxes


def _add_injection(shapes, x: float, lane_y: float, number: int, note: str, drop: float = 0.0) -> None:
    """Mark an injection checkpoint; ``drop`` lowers the note below a neighbouring one."""
    marker = shapes.add_group_shape()
    marker.name = f"Injection {number}"
    top, bottom = lane_y - STEP_H / 2 - 0.08, lane_y + STEP_H / 2 + 0.08 + drop
    _red_line(marker.shapes, x, top, bottom, "Injection line")
    _badge(marker.shapes, x, top - 0.16, number, "Injection badge", radius=0.15, size_pt=12)
    _text(marker.shapes, x - 1.0, bottom + 0.03, 2.0, 0.26, "Injection note", note, size_pt=11, color=FAILURE,
          align=PP_ALIGN.CENTER)
    marker._element.recalculate_extents()


def build_timeline() -> Presentation:
    """Build the injection timeline figure.

    Returns:
        A one-slide presentation with the KET4_Square_4mm and gear_small routes.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(TIMELINE_HEIGHT_IN)
    shapes = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes

    # Markers go first so the step boxes cover their lines; only the cut above and below shows.
    _add_injection(shapes, _step_left(1) + 0.18 * STEP_W, KET_LANE_Y, 3, "18% trimmed")
    _add_injection(shapes, _step_left(2) + 0.55 * STEP_W, KET_LANE_Y, 2, "before grasp")
    _add_injection(shapes, _step_left(3) - STEP_GAP / 2, KET_LANE_Y, 1, "held, before release", drop=0.3)
    _add_injection(shapes, _step_left(5) + 0.5 * STEP_W, GEAR_LANE_Y, 4, "50% lowered")

    _add_lane_label(shapes, KET_LANE_Y, "KET", "KET4_Square_4mm", "trimmed at M1")
    _add_lane_label(shapes, GEAR_LANE_Y, "gear", "gear_small", "printed")
    ket_boxes = _add_lane(shapes, KET_LANE_Y, KET_STEPS)
    gear_boxes = _add_lane(shapes, GEAR_LANE_Y, GEAR_STEPS)

    board_top = KET_LANE_Y - STEP_H / 2
    board = _resource_box(shapes, _step_left(6), board_top, STEP_W, GEAR_LANE_Y + STEP_H / 2 - board_top,
                          "Assembly Board", "product", [("Assembly", 14, True, INK), ("Board", 14, True, INK)])
    board.line.width = Pt(2.25)
    _arrow(shapes, ket_boxes[-1], board, "Flow UR5e-3 assemble to Assembly Board", MSO_CONNECTOR.ELBOW)
    _arrow(shapes, gear_boxes[-1], board, "Flow UR5e-4 place to Assembly Board", MSO_CONNECTOR.ELBOW)

    _text(shapes, MARGIN, TIMELINE_HEIGHT_IN - 0.42, 6.0, 0.3, "Note",
          "Each failure is injected once, in its own run.", size_pt=11, color=MUTED)
    return presentation


def main() -> None:
    """Crop the snapshots and write both failure figures next to this script."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    _crop_snapshots()
    for build, path in ((build_panels, PANELS_PATH), (build_timeline, TIMELINE_PATH)):
        build().save(path)
        logger.info("Wrote %s", path)


if __name__ == "__main__":
    main()
