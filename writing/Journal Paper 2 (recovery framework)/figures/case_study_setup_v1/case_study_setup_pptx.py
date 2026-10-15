"""Write the case-study cell schematic as editable PowerPoint shapes.

This draws the same top view as ``case_study_setup.py`` using native shapes,
so stations, labels, and arrows can be moved and restyled in PowerPoint.
Each station, robot, and legend is a named group (see the Selection Pane).

Slide 1 shows the layout and part flow; slide 2 adds the four failure
scenarios. Positions are in metres from the configured Gazebo layout and are
drawn at 1.48 in per metre on a 13.333 x 7.5 in slide.

Requires ``python-pptx`` (not a project dependency):

    python case_study_setup_pptx.py
"""

from __future__ import annotations

import logging
import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

logger = logging.getLogger(__name__)

OUTPUT_PATH = Path(__file__).resolve().parent / "case_study_setup.pptx"
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5
SCALE_IN_PER_M = 1.48
X_LIMITS = (-7.10, 1.72)
Y_LIMITS = (-1.42, 3.42)
OFFSET_X_IN = (SLIDE_WIDTH_IN - (X_LIMITS[1] - X_LIMITS[0]) * SCALE_IN_PER_M) / 2
OFFSET_Y_IN = (SLIDE_HEIGHT_IN - (Y_LIMITS[1] - Y_LIMITS[0]) * SCALE_IN_PER_M) / 2
FONT = "Arial"

INK = "2B2F36"
MUTED = "5D6673"
FLOW = "1F5FA8"
KMR_ROUTE = "B8661D"
KMR_BODY = "F0A648"
KMR_EDGE = "8A5A14"
ROBOT = "2F3B4C"
REACH = "5B7BA6"
MACHINE = "DDE1E7"
STORAGE = "DCE8F5"
STORAGE_EDGE = "3D5A80"
TABLE = "F6E9D4"
TABLE_EDGE = "B08850"
BELT = "6B7280"
ROLLER = "8B929C"
LOADING = "3BB39B"
STAGING = "BFE3DC"
STAGING_EDGE = "2A8C7A"
PEG = "4A78B5"
GEAR = "E0843C"
PRINTER = "4B5563"
PRINTER_BED = "A9B0BA"
FAILURE = "D62728"
WHITE = "FFFFFF"

LABEL_PT = 14
MACHINE_LABEL_PT = 15
SMALL_PT = 12
BADGE_PT = 11
OUTLINE_PT = 1.5
THIN_PT = 1.0
FLOW_PT = 2.0
ARM_PT = 4.4

FAILURE_SCENARIOS = (
    "Conveyor breakdown",
    "ur5e-1 breakdown",
    "Machining breakdown during part processing",
    "Part slippage",
)

Point = tuple[float, float]


def _x(x_m: float) -> Emu:
    return Inches(OFFSET_X_IN + (x_m - X_LIMITS[0]) * SCALE_IN_PER_M)


def _y(y_m: float) -> Emu:
    return Inches(OFFSET_Y_IN + (Y_LIMITS[1] - y_m) * SCALE_IN_PER_M)


def _length(m: float) -> Emu:
    return Inches(m * SCALE_IN_PER_M)


def _style_line(shape, color: str | None, width_pt: float = THIN_PT, dash=None) -> None:
    if color is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width_pt)
    if dash is not None:
        shape.line.dash_style = dash


def _style_fill(shape, color: str | None, alpha_percent: int | None = None) -> None:
    if color is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    if alpha_percent is not None:
        srgb = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(srgb, qn("a:alpha")).set("val", str(alpha_percent * 1000))


def _finish_line(shape, *, arrow_end: bool = False, round_cap: bool = False) -> None:
    ln = shape.line._get_or_add_ln()
    if round_cap:
        ln.set("cap", "rnd")
        etree.SubElement(ln, qn("a:round"))
    if arrow_end:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def _box(shapes, center: Point, size: Point, name: str, *, fill: str | None, line: str | None = INK,
         line_pt: float = OUTLINE_PT, dash=None, shape_type=MSO_SHAPE.RECTANGLE, corner_m: float = 0.0,
         alpha_percent: int | None = None):
    (cx, cy), (w, h) = center, size
    shape = shapes.add_shape(shape_type, _x(cx - w / 2), _y(cy + h / 2), _length(w), _length(h))
    shape.name = name
    shape.shadow.inherit = False
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = corner_m / min(w, h)
    _style_fill(shape, fill, alpha_percent)
    _style_line(shape, line, line_pt, dash)
    return shape


def _circle(shapes, center: Point, radius: float, name: str, *, fill: str | None, line: str | None = None,
            line_pt: float = THIN_PT, dash=None, alpha_percent: int | None = None):
    return _box(shapes, center, (2 * radius, 2 * radius), name, fill=fill, line=line, line_pt=line_pt,
                dash=dash, shape_type=MSO_SHAPE.OVAL, alpha_percent=alpha_percent)


def _segment(shapes, start: Point, end: Point, name: str, *, color: str, width_pt: float = THIN_PT,
             dash=None, arrow_end: bool = False):
    connector = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _x(start[0]), _y(start[1]), _x(end[0]), _y(end[1]))
    connector.name = name
    _style_line(connector, color, width_pt, dash)
    _finish_line(connector, arrow_end=arrow_end)
    return connector


def _polyline(shapes, points: list[Point], name: str, *, color: str, width_pt: float, dash=None,
              arrow_end: bool = False, round_cap: bool = False, curved: bool = False):
    """Add an open freeform; with ``curved`` the three points are start, control, end."""
    builder = shapes.build_freeform(_x(points[0][0]), _y(points[0][1]))
    builder.add_line_segments([(_x(x), _y(y)) for x, y in points[1:]], close=False)
    shape = builder.convert_to_shape()
    shape.name = name
    shape.shadow.inherit = False
    path = shape._element.spPr.find(qn("a:custGeom")).find(qn("a:pathLst")).find(qn("a:path"))
    if curved:
        line_tos = path.findall(qn("a:lnTo"))
        curve = etree.SubElement(path, qn("a:quadBezTo"))
        for line_to in line_tos:
            curve.append(line_to.find(qn("a:pt")))
            path.remove(line_to)
    path.set("fill", "none")
    _style_fill(shape, None)
    _style_line(shape, color, width_pt, dash)
    _finish_line(shape, arrow_end=arrow_end, round_cap=round_cap)
    return shape


def _text(shapes, position: Point, text: str, name: str, *, size_pt: float = LABEL_PT, bold: bool = True,
          color: str = INK, align: str = "center", anchor: str = "middle", width_m: float = 1.6):
    """Add a label; ``position`` is the left, centre, or top reference of the text."""
    lines = text.split("\n")
    height = Pt(size_pt * 1.25 * len(lines))
    width = _length(width_m)
    left = _x(position[0]) if align == "left" else _x(position[0]) - width // 2
    top = _y(position[1]) if anchor == "top" else _y(position[1]) - height // 2
    box = shapes.add_textbox(left, top, width, height)
    box.name = name
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP if anchor == "top" else MSO_ANCHOR.MIDDLE
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT if align == "left" else PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return box


def _group(shapes, name: str):
    group = shapes.add_group_shape()
    group.name = name
    return group


def _close_group(group) -> None:
    group._element.recalculate_extents()


def _square_peg(shapes, center: Point, name: str, size: float = 0.085) -> None:
    _box(shapes, center, (size, size), name, fill=PEG, line=WHITE, line_pt=0.75)


def _round_peg(shapes, center: Point, name: str, radius: float = 0.048) -> None:
    _circle(shapes, center, radius, name, fill=PEG, line=WHITE, line_pt=0.75)


def _gear(shapes, center: Point, radius: float, name: str) -> None:
    _box(shapes, center, (2 * radius, 2 * radius), name, fill=GEAR, line=None, shape_type=MSO_SHAPE.GEAR_9)


def _quadratic_flow_points(start: Point, end: Point, rad: float) -> list[Point]:
    # Same control point as matplotlib's arc3 connection style.
    (x1, y1), (x2, y2) = start, end
    dx, dy = x2 - x1, y2 - y1
    control = ((x1 + x2) / 2 + rad * dy, (y1 + y2) / 2 - rad * dx)
    return [start, control, end]


def _flow(shapes, start: Point, end: Point, name: str, rad: float = 0.0) -> None:
    if rad == 0.0:
        _segment(shapes, start, end, name, color=FLOW, width_pt=FLOW_PT, arrow_end=True)
        return
    _polyline(shapes, _quadratic_flow_points(start, end, rad), name, color=FLOW, width_pt=FLOW_PT,
              arrow_end=True, curved=True)


def _kmr_route(shapes, points: list[Point], name: str) -> None:
    if len(points) == 2:
        _segment(shapes, points[0], points[1], name, color=KMR_ROUTE, width_pt=FLOW_PT,
                 dash=MSO_LINE_DASH_STYLE.DASH, arrow_end=True)
        return
    _polyline(shapes, points, name, color=KMR_ROUTE, width_pt=FLOW_PT, dash=MSO_LINE_DASH_STYLE.DASH,
              arrow_end=True)


def _reach(shapes, center: Point, name: str) -> None:
    _circle(shapes, center, 0.85, name, fill=REACH, line=REACH, line_pt=THIN_PT,
            dash=MSO_LINE_DASH_STYLE.DASH, alpha_percent=8)


def _robot(shapes, center: Point, heading_deg: float, name: str) -> None:
    """Add a UR5e as a base disc with a folded two-link arm pointing at its work."""
    x, y = center
    heading = math.radians(heading_deg)
    elbow = (x + 0.22 * math.cos(heading + 0.55), y + 0.22 * math.sin(heading + 0.55))
    wrist = (elbow[0] + 0.20 * math.cos(heading - 0.55), elbow[1] + 0.20 * math.sin(heading - 0.55))
    finger = (-math.sin(heading) * 0.045, math.cos(heading) * 0.045)
    tip = (math.cos(heading) * 0.06, math.sin(heading) * 0.06)

    group = _group(shapes, name)
    _polyline(group.shapes, [center, elbow, wrist], f"{name} arm", color=ROBOT, width_pt=ARM_PT, round_cap=True)
    gripper = [
        (wrist[0] + finger[0] + tip[0], wrist[1] + finger[1] + tip[1]),
        (wrist[0] + finger[0], wrist[1] + finger[1]),
        (wrist[0] - finger[0], wrist[1] - finger[1]),
        (wrist[0] - finger[0] + tip[0], wrist[1] - finger[1] + tip[1]),
    ]
    _polyline(group.shapes, gripper, f"{name} gripper", color=ROBOT, width_pt=2.0, round_cap=True)
    _circle(group.shapes, center, 0.12, f"{name} base", fill=ROBOT, line=WHITE, line_pt=1.1)
    _circle(group.shapes, center, 0.045, f"{name} base centre", fill=WHITE)
    _close_group(group)


def _failure_badge(shapes, center: Point, number: int, name: str, radius: float = 0.105) -> None:
    badge = _circle(shapes, center, radius, name, fill=FAILURE, line=WHITE, line_pt=1.5)
    frame = badge.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    run.font.name = FONT
    run.font.size = Pt(BADGE_PT)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)


def _draw_background(shapes) -> None:
    for name, center in (("UR5e-1", (-4.05, 1.10)), ("UR5e-2", (-2.05, 1.10)),
                         ("UR5e-3", (0.0, 0.50)), ("UR5e-4", (0.0, -0.50))):
        _reach(shapes, center, f"{name} reach")
    _box(shapes, (0.0, 0.0), (1.50, 1.60), "Assembly table", fill=TABLE, line=TABLE_EDGE)


def _draw_storage_and_kmr(shapes) -> None:
    storage = _group(shapes, "Storage")
    _box(storage.shapes, (-6.45, 2.30), (1.00, 1.25), "Storage body", fill=STORAGE, line=STORAGE_EDGE)
    for tray_x, tray_name in ((-6.62, "KET tray"), (-6.28, "RGOCG tray")):
        _box(storage.shapes, (tray_x, 2.30), (0.22, 1.08), tray_name, fill=WHITE, line=STORAGE_EDGE, line_pt=THIN_PT)
    for slot, slot_y in enumerate((1.92, 2.17, 2.42, 2.67), start=1):
        _square_peg(storage.shapes, (-6.62, slot_y), f"KET peg {slot}")
        _round_peg(storage.shapes, (-6.28, slot_y), f"RGOCG peg {slot}")
    _close_group(storage)

    # KMR docking poses; the platform stands sideways (yaw +90 deg) at every dock.
    for dock_x, machine in ((-4.70, "M1"), (-2.70, "M2")):
        _box(shapes, (dock_x, 2.20), (0.63, 1.08), f"KMR docking pose {machine}", fill=None, line=KMR_ROUTE,
             line_pt=1.25, dash=MSO_LINE_DASH_STYLE.SQUARE_DOT, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
             corner_m=0.06)
    _kmr_route(shapes, [(-5.54, 2.13), (-4.70, 2.20)], "KMR route Storage-M1")
    _kmr_route(shapes, [(-5.54, 2.13), (-5.54, 3.10), (-2.70, 3.10), (-2.70, 2.20)], "KMR route Storage-M2")

    kmr = _group(shapes, "KUKA KMR")
    _box(kmr.shapes, (-5.54, 2.13), (0.63, 1.08), "KMR platform", fill=KMR_BODY, line=KMR_EDGE,
         shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, corner_m=0.06)
    for wheel_x in (-5.80, -5.28):
        for wheel_y in (1.74, 2.52):
            _box(kmr.shapes, (wheel_x, wheel_y), (0.07, 0.22), "KMR wheel", fill=INK, line=None)
    _robot(kmr.shapes, (-5.54, 1.88), 165.0, "KMR arm")
    _close_group(kmr)

    _text(shapes, (-6.45, 1.53), "Storage", "Label Storage", anchor="top")
    _text(shapes, (-5.54, 1.53), "KUKA KMR", "Label KUKA KMR", anchor="top")


def _draw_machine_cell(shapes, name: str, x: float, robot_name: str, peg_shape: str) -> None:
    machine = _group(shapes, name)
    _box(machine.shapes, (x, 1.90), (0.531, 0.503), f"{name} enclosure", fill=MACHINE, line=INK)
    _text(machine.shapes, (x, 1.97), name, f"{name} name", size_pt=MACHINE_LABEL_PT, width_m=0.5)
    if peg_shape == "square":
        _square_peg(machine.shapes, (x, 1.79), f"{name} peg", size=0.075)
    else:
        _round_peg(machine.shapes, (x, 1.79), f"{name} peg", radius=0.042)
    _close_group(machine)

    _box(shapes, (x - 0.55, 1.10), (0.16, 0.16), f"{name} staging tray", fill=STAGING, line=STAGING_EDGE,
         line_pt=THIN_PT)
    _robot(shapes, (x, 1.10), 90.0, robot_name)
    _text(shapes, (x + 0.19, 1.06), robot_name, f"Label {robot_name}", align="left")


def _draw_conveyor(shapes) -> None:
    conveyor = _group(shapes, "Conveyor")
    _box(conveyor.shapes, (-2.85, 0.50), (4.20, 0.30), "Conveyor belt", fill=BELT, line=INK)
    roller_x = -4.85
    while roller_x < -0.80:
        _segment(conveyor.shapes, (roller_x, 0.36), (roller_x, 0.64), "Conveyor roller", color=ROLLER, width_pt=0.75)
        roller_x += 0.15
    for chevron_x in (-4.55, -3.55, -2.55, -1.55, -1.05):
        _polyline(conveyor.shapes, [(chevron_x - 0.04, 0.565), (chevron_x + 0.04, 0.50), (chevron_x - 0.04, 0.435)],
                  "Conveyor direction", color=WHITE, width_pt=1.6)
    for loading_x, robot_name in ((-4.05, "ur5e-1"), (-2.05, "ur5e-2")):
        _box(conveyor.shapes, (loading_x, 0.50), (0.20, 0.22), f"Conveyor loading position {robot_name}",
             fill=LOADING, line=WHITE, line_pt=THIN_PT)
    _close_group(conveyor)
    _text(shapes, (-3.05, 0.27), "Conveyor", "Label Conveyor", anchor="top")


def _draw_assembly_station(shapes) -> None:
    buffer = _group(shapes, "Buffer")
    zone_edges = (-0.74, -0.62, -0.50, -0.38, -0.26)
    for zone, (left, right) in enumerate(zip(zone_edges, zone_edges[1:]), start=1):
        _box(buffer.shapes, ((left + right) / 2, 0.50), (right - left, 0.24), f"Buffer zone {zone}",
             fill="9FBCDC" if zone == 4 else "CFE0F1", line=STORAGE_EDGE, line_pt=THIN_PT)
    _close_group(buffer)

    board = _group(shapes, "Assembly Board")
    _box(board.shapes, (0.0, 0.0), (0.384, 0.384), "Assembly Board plate", fill=WHITE, line=INK)
    for column, hole_x in enumerate((-0.135, -0.045, 0.045, 0.135)):
        _square_peg(board.shapes, (hole_x, 0.11), "Board KET peg", size=0.05)
        _round_peg(board.shapes, (hole_x, 0.025), "Board RGOCG peg", radius=0.028)
        if column < 3:
            _gear(board.shapes, (-0.09 + column * 0.09, -0.10), (0.028, 0.035, 0.042)[column], "Board gear")
    _close_group(board)

    _box(shapes, (0.50, 0.58), (0.24, 0.18), "Exit", fill=STORAGE, line=STORAGE_EDGE, line_pt=1.25)

    printer = _group(shapes, "Printer")
    _box(printer.shapes, (0.50, -0.50), (0.36, 0.46), "Printer frame", fill=PRINTER, line=INK)
    _box(printer.shapes, (0.45, -0.50), (0.18, 0.34), "Printer bed", fill=PRINTER_BED, line=None)
    for gear_y, radius, gear_name in ((-0.58, 0.028, "gear_small"), (-0.50, 0.035, "gear_medium"),
                                      (-0.42, 0.042, "gear_large")):
        _gear(printer.shapes, (0.44, gear_y), radius, gear_name)
    _close_group(printer)

    _robot(shapes, (0.0, 0.50), 180.0, "UR5e-3")
    _robot(shapes, (0.0, -0.50), 0.0, "UR5e-4")

    _flow(shapes, (-0.30, 0.37), (-0.15, 0.21), "Flow Buffer to Assembly Board", rad=0.25)
    _flow(shapes, (0.31, -0.40), (0.20, -0.14), "Flow Printer to Assembly Board", rad=-0.25)
    _flow(shapes, (0.20, 0.15), (0.37, 0.50), "Flow Assembly Board to Exit", rad=0.30)

    _segment(shapes, (0.20, 0.0), (0.80, 0.0), "Assembly Board leader", color=MUTED, width_pt=THIN_PT)
    _text(shapes, (-0.50, 0.98), "Buffer", "Label Buffer")
    _text(shapes, (0.10, 0.98), "UR5e-3", "Label UR5e-3", align="left")
    _text(shapes, (0.0, -0.86), "UR5e-4", "Label UR5e-4", anchor="top")
    _text(shapes, (0.84, 0.58), "Exit", "Label Exit", align="left")
    _text(shapes, (0.84, 0.0), "Assembly\nBoard", "Label Assembly Board", align="left")
    _text(shapes, (0.84, -0.50), "Printer", "Label Printer", align="left")


def _draw_legend(shapes) -> None:
    legend = _group(shapes, "Legend")
    group_shapes = legend.shapes
    left, top, step = -7.02, 1.12, 0.215
    symbol_x = left + 0.13
    rows = (
        "Part transfer",
        "KMR route",
        "KMR docking pose",
        "Conveyor loading position",
        "Staging tray",
        "UR5e reach (0.85 m)",
        "KET square peg (×4)",
        "RGOCG round peg (×4)",
        "Printed gear (×3)",
    )
    for index, text in enumerate(rows):
        y = top - index * step
        _text(group_shapes, (left + 0.34, y), text, f"Legend text {text}", size_pt=SMALL_PT, bold=False,
              align="left", width_m=1.65)
        if text == "Part transfer":
            _segment(group_shapes, (left, y), (left + 0.26, y), "Legend part transfer", color=FLOW,
                     width_pt=FLOW_PT, arrow_end=True)
        elif text == "KMR route":
            _segment(group_shapes, (left, y), (left + 0.26, y), "Legend KMR route", color=KMR_ROUTE,
                     width_pt=FLOW_PT, dash=MSO_LINE_DASH_STYLE.DASH, arrow_end=True)
        elif text == "KMR docking pose":
            _box(group_shapes, (symbol_x, y), (0.14, 0.14), "Legend KMR docking pose", fill=None, line=KMR_ROUTE,
                 line_pt=1.25, dash=MSO_LINE_DASH_STYLE.SQUARE_DOT, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
                 corner_m=0.02)
        elif text == "Conveyor loading position":
            _box(group_shapes, (symbol_x, y), (0.20, 0.13), "Legend loading position", fill=LOADING, line=None)
        elif text == "Staging tray":
            _box(group_shapes, (symbol_x, y), (0.13, 0.13), "Legend staging tray", fill=STAGING,
                 line=STAGING_EDGE, line_pt=THIN_PT)
        elif text.startswith("UR5e reach"):
            _circle(group_shapes, (symbol_x, y), 0.09, "Legend reach", fill=REACH, line=REACH, line_pt=THIN_PT,
                    dash=MSO_LINE_DASH_STYLE.DASH, alpha_percent=14)
        elif text.startswith("KET"):
            _square_peg(group_shapes, (symbol_x, y), "Legend KET peg")
        elif text.startswith("RGOCG"):
            _round_peg(group_shapes, (symbol_x, y), "Legend RGOCG peg")
        else:
            _gear(group_shapes, (symbol_x, y), 0.065, "Legend gear")
    _close_group(legend)


def _draw_scale_bar(shapes) -> None:
    scale_bar = _group(shapes, "Scale bar")
    x0, y0 = -4.90, -0.80
    _segment(scale_bar.shapes, (x0, y0), (x0 + 1.0, y0), "Scale bar line", color=INK, width_pt=1.75)
    for tick_x in (x0, x0 + 0.5, x0 + 1.0):
        _segment(scale_bar.shapes, (tick_x, y0 - 0.04), (tick_x, y0 + 0.04), "Scale bar tick", color=INK,
                 width_pt=1.25)
    _text(scale_bar.shapes, (x0 + 0.5, y0 + 0.13), "1 m", "Scale bar label", size_pt=SMALL_PT, bold=False,
          width_m=0.6)
    _close_group(scale_bar)


def _draw_failures(shapes) -> None:
    markers = _group(shapes, "Failure markers")
    _failure_badge(markers.shapes, (-3.05, 0.50), 1, "Failure 1 Conveyor breakdown")
    _failure_badge(markers.shapes, (-3.80, 1.38), 2, "Failure 2 ur5e-1 breakdown")
    _failure_badge(markers.shapes, (-3.79, 2.15), 3, "Failure 3 Machining breakdown during part processing")
    # Part slippage: ur5e-4 releases gear_small early; it rests upright in ur5e-3's region.
    _failure_badge(markers.shapes, (0.0, 0.287), 4, "Failure 4 Part slippage", radius=0.092)
    _close_group(markers)

    failure_legend = _group(shapes, "Failure scenarios legend")
    left, top, step = -1.50, 3.28, 0.235
    _text(failure_legend.shapes, (left, top), "Failure scenarios", "Failure scenarios title", size_pt=13,
          align="left", width_m=2.0)
    for number, scenario in enumerate(FAILURE_SCENARIOS, start=1):
        y = top - number * step
        _failure_badge(failure_legend.shapes, (left + 0.09, y), number, f"Failure legend badge {number}",
                       radius=0.085)
        _text(failure_legend.shapes, (left + 0.26, y), scenario, f"Failure legend text {number}",
              size_pt=SMALL_PT, bold=False, align="left", width_m=2.9)
    _close_group(failure_legend)


def _draw_slide(shapes, show_failures: bool) -> None:
    _draw_background(shapes)
    _draw_storage_and_kmr(shapes)
    _draw_machine_cell(shapes, "M1", -4.05, "UR5e-1", "square")
    _draw_machine_cell(shapes, "M2", -2.05, "UR5e-2", "round")
    _draw_conveyor(shapes)
    _flow(shapes, (-4.13, 1.635), (-4.13, 0.63), "Flow M1 to Conveyor", rad=0.32)
    _flow(shapes, (-2.13, 1.635), (-2.13, 0.63), "Flow M2 to Conveyor", rad=0.32)
    _draw_assembly_station(shapes)
    _draw_legend(shapes)
    _draw_scale_bar(shapes)
    if show_failures:
        _draw_failures(shapes)


def build_presentation() -> Presentation:
    """Build the two-slide schematic deck.

    Returns:
        A presentation with the layout slide and the failure-scenario slide.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = presentation.slide_layouts[6]
    for show_failures in (False, True):
        slide = presentation.slides.add_slide(blank_layout)
        _draw_slide(slide.shapes, show_failures)
    return presentation


def main() -> None:
    """Write ``case_study_setup.pptx`` next to this script."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    build_presentation().save(OUTPUT_PATH)
    logger.info("Wrote %s", OUTPUT_PATH)


if __name__ == "__main__":
    main()
