"""Write the Part slippage recovery outline as one editable PowerPoint slide.

The slide draws the outline of ``draft_66d30ad65cc84e959dc726686fd62ae3``
(gpt-5.4, reasoning effort medium, every check off, 20 Oct 2026) as a
resource-lane diagram: the fault, the four recovery events the LLM proposed,
the neighbor state one event changes, the handoff, and the nominal events that
connect to the remaining plan. Review findings come from the 20 Oct entry in
``IMPLEMENTATION_PLAN.md``. Colours, fonts, and resource categories match
``case_study_failures_v1``.

Requires ``python-pptx`` (not a project dependency):

    python part_slippage_recovery_outline_pptx.py
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

FIGURE_DIR = Path(__file__).resolve().parent
SNAPSHOT_PATH = FIGURE_DIR.parent / "case_study_failures_v1" / "failure_4.png"
OUTPUT_PATH = FIGURE_DIR / "part_slippage_recovery_outline.pptx"

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5
MARGIN = 0.5
FONT = "Arial"

INK = "2B2F36"
MUTED = "5D6673"
FLOW = "1F5FA8"
FAILURE = "D62728"
WHITE = "FFFFFF"
PANEL_EDGE = "D5DAE1"
CARD_FILL = "F4F6F8"
LANE_BAND = "F7F8FA"
EVENT_FILL = "E4EEFB"
NOMINAL_EDGE = "9AA3AE"
REVIEW = "F2A900"

# (fill, border) for each resource category, as in the case study figures.
CATEGORY_STYLE = {
    "storage": ("DCE8F5", "3D5A80"),
    "robot": ("D7EEE7", "2A8C7A"),
}

# Connection-site indices of a rectangle preset.
LEFT, RIGHT = 1, 3

# Diagram grid, in inches.
HEADER_Y = 1.86
LANE_TOP = 2.14
LANE_H = 0.8
LANE_LABEL_X = MARGIN
FAULT_X = 1.55
FAULT_W = 0.7
EVENT_X0 = FAULT_X + FAULT_W
EVENT_COL_W = 1.38
EVENT_W = 1.26
EVENT_H = 0.62
HANDOFF_W = 0.76
NOMINAL_COL_W = 1.16
NOMINAL_W = 1.04


@dataclass(frozen=True)
class Event:
    """One box in the diagram: a recovery event or a nominal connection event."""

    lane: str
    title: str
    detail: str
    name: str
    # (lane, change) of an adjacent resource that takes no action but whose state the event changes.
    shared: tuple[str, str] | None = None


LANES = (("ur5e-4", "robot"), ("ur5e-3", "robot"), ("Buffer", "storage"))
FAULT_LANE = "ur5e-4"

RECOVERY_EVENTS = (
    Event("ur5e-4", "retract home", "clear fault", "clear_fault_and_retract_home"),
    Event("ur5e-3", "place KET4", "in buffer zone_1", "place_KET4_in_buffer_zone_1", shared=("Buffer", "zone_1 ← KET4")),
    Event("ur5e-3", "pick gear_small", "from drop pose", "recover_pick_gear_small_from_drop_pose"),
    Event("ur5e-3", "assemble gear_small", "on assembly board", "assemble_place_gear_small_on_assembly_board_target"),
)
NOMINAL_EVENTS = (
    Event("ur5e-3", "move_home", "", "move_home"),
    Event("ur5e-3", "pick_approach", "KET4", "pick_approach"),
    Event("ur5e-3", "pick_grasp", "KET4", "pick_grasp"),
)
FLAGGED_STEP = 2

SPEAKER_NOTES = """\
Part slippage: ur5e-4 dropped gear_small during place_approach (device fault report from gripper object detection; measured holder: none).

The LLM (gpt-5.4, reasoning effort medium, every check off) proposed four recovery events that do not exist in the nominal DES models:
1. ur5e-4 clear_fault_and_retract_home
2. ur5e-3 place_KET4_in_buffer_zone_1 (also changes the Buffer: zone_1_part = KET4_Square_4mm)
3. ur5e-3 recover_pick_gear_small_from_drop_pose
4. ur5e-3 assemble_place_gear_small_on_assembly_board_target

Code then found nominal events that connect to the remaining plan: ur5e-3 move_home, pick_approach (KET4), pick_grasp (KET4). The handoff check passed.

Review findings: the handoff holds only on ur5e-3's reduced pick model; the buffer's own pick guard needs KET4 in zone_4_part, but the LLM put it in zone_1. gear_small's process record never gets the assembly that safety rule 5a reads.

Run: draft_66d30ad65cc84e959dc726686fd62ae3, 8 LLM calls, 243 s, about 82k tokens."""


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _style_run(run, size_pt: float, *, bold: bool = False, color: str = INK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _set_margins(frame, inches: float) -> None:
    frame.margin_left = frame.margin_right = Inches(inches)
    frame.margin_top = frame.margin_bottom = Inches(inches / 2)


def _write_lines(frame, lines, *, align=PP_ALIGN.LEFT, space_after_pt: float = 0) -> None:
    """Fill a text frame with (text, size, bold, color) lines, one paragraph each."""
    frame.word_wrap = True
    for index, (text, size_pt, bold, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after_pt)
        run = paragraph.add_run()
        run.text = text
        _style_run(run, size_pt, bold=bold, color=color)


def _add_text(shapes, x, y, w, h, lines, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after_pt: float = 0,
              name: str = ""):
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name or box.name
    frame = box.text_frame
    _set_margins(frame, 0)
    frame.vertical_anchor = anchor
    _write_lines(frame, lines, align=align, space_after_pt=space_after_pt)
    return box


def _add_shape(shapes, kind, x, y, w, h, *, fill: str | None, line: str | None, line_pt: float = 1.0,
               dashed: bool = False, name: str = ""):
    shape = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or shape.name
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_pt)
        if dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def _add_arrowhead(connector) -> None:
    line = connector.line._get_or_add_ln()
    tail = etree.SubElement(line, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _connect(shapes, begin, begin_site, end, end_site):
    connector = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    connector.begin_connect(begin, begin_site)
    connector.end_connect(end, end_site)
    connector.line.color.rgb = _rgb(MUTED)
    connector.line.width = Pt(1.25)
    _add_arrowhead(connector)
    return connector


def _line(shapes, x1, y1, x2, y2, *, color: str = MUTED, width_pt: float = 1.25, arrow: bool = True,
          dashed: bool = False):
    connector = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = _rgb(color)
    connector.line.width = Pt(width_pt)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        _add_arrowhead(connector)
    return connector


def _lane_center(lane: str) -> float:
    return LANE_TOP + _lane_index(lane) * LANE_H + LANE_H / 2


def _event_x(step: int) -> float:
    return EVENT_X0 + (step - 1) * EVENT_COL_W + (EVENT_COL_W - EVENT_W) / 2


def _handoff_x() -> float:
    return EVENT_X0 + len(RECOVERY_EVENTS) * EVENT_COL_W + HANDOFF_W / 2


def _nominal_x(index: int) -> float:
    start = EVENT_X0 + len(RECOVERY_EVENTS) * EVENT_COL_W + HANDOFF_W
    return start + index * NOMINAL_COL_W + (NOMINAL_COL_W - NOMINAL_W) / 2


def _lane_index(lane: str) -> int:
    return [name for name, _ in LANES].index(lane)


def _add_event_box(shapes, x: float, event: Event, *, nominal: bool, label: str):
    centers = [_lane_center(event.lane)]
    if event.shared:
        centers.append(_lane_center(event.shared[0]))
    top = min(centers) - EVENT_H / 2
    height = max(centers) - min(centers) + EVENT_H
    box = _add_shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, top, NOMINAL_W if nominal else EVENT_W, height,
                     fill=CARD_FILL if nominal else EVENT_FILL, line=NOMINAL_EDGE if nominal else FLOW,
                     line_pt=1.0 if nominal else 1.5, name=label)
    box.adjustments[0] = 0.12 * EVENT_H / height
    frame = box.text_frame
    _set_margins(frame, 0.06)
    # A box that spans a second lane keeps its own text in the acting resource's lane.
    if event.shared is None:
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif centers[0] < centers[1]:
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_top = Inches(0.1)
    else:
        frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        frame.margin_bottom = Inches(0.1)
    lines = [(event.title, 10 if nominal else 10.5, False, MUTED if nominal else INK)]
    if event.detail:
        lines.append((event.detail, 9, False, MUTED))
    _write_lines(frame, lines, align=PP_ALIGN.CENTER)
    if event.shared:
        _add_shared_change(shapes, x, event)
    return box


def _add_shared_change(shapes, x: float, event: Event) -> None:
    """Show, inside the event box, the state change of the adjacent resource that takes no action."""
    lane, change = event.shared
    boundary_y = LANE_TOP + max(_lane_index(lane), _lane_index(event.lane)) * LANE_H
    _line(shapes, x + 0.1, boundary_y, x + EVENT_W - 0.1, boundary_y, color=FLOW, width_pt=0.75, arrow=False,
          dashed=True)
    center = _lane_center(lane)
    _add_text(shapes, x, center - 0.15, EVENT_W, 0.3, [(change, 10.5, False, FLOW)], align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, name=f"{lane} change")


def _add_review_flag(shapes, x: float, y: float, name: str) -> None:
    flag = _add_shape(shapes, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, 0.28, 0.24, fill=REVIEW, line=WHITE,
                      line_pt=1.0, name=name)
    frame = flag.text_frame
    _set_margins(frame, 0)
    frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    _write_lines(frame, [("!", 9, True, INK)], align=PP_ALIGN.CENTER)


def _draw_header(shapes) -> None:
    _add_text(shapes, MARGIN, 0.36, SLIDE_WIDTH_IN - 2 * MARGIN, 0.6,
              [("Part slippage: recovery outline proposed by the LLM", 28, True, INK)], name="Title")
    _add_text(shapes, MARGIN, 0.98, SLIDE_WIDTH_IN - 2 * MARGIN, 0.32,
              [("gpt-5.4, reasoning effort medium, every check off  ·  run of 20 Oct 2026  ·  "
                "4 recovery events, 8 LLM calls, 243 s", 13, False, MUTED)], name="Subtitle")

    legend_y = 1.42
    x = MARGIN
    items = (
        ("fault", "fault"),
        ("event", "new recovery event (not in the nominal models)"),
        ("nominal", "nominal event"),
        ("flag", "issue found in review"),
    )
    for kind, text in items:
        if kind == "fault":
            _add_shape(shapes, MSO_SHAPE.OVAL, x, legend_y + 0.04, 0.18, 0.18, fill=FAILURE, line=None, name="Legend fault")
        elif kind == "flag":
            _add_shape(shapes, MSO_SHAPE.ISOSCELES_TRIANGLE, x, legend_y + 0.03, 0.2, 0.18, fill=REVIEW, line=None,
                       name="Legend flag")
        else:
            fill, line, dashed = {
                "event": (EVENT_FILL, FLOW, False),
                "nominal": (CARD_FILL, NOMINAL_EDGE, False),
            }[kind]
            _add_shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, legend_y + 0.03, 0.3, 0.2, fill=fill, line=line,
                       dashed=dashed, name=f"Legend {kind}")
        width = 0.068 * len(text) + 0.1
        _add_text(shapes, x + 0.38, legend_y, width, 0.26, [(text, 11, False, MUTED)],
                  anchor=MSO_ANCHOR.MIDDLE, name=f"Legend {kind} text")
        x += 0.38 + width + 0.22


def _draw_lanes(shapes) -> None:
    for index, (lane, category) in enumerate(LANES):
        top = LANE_TOP + index * LANE_H
        if index % 2 == 0:
            _add_shape(shapes, MSO_SHAPE.RECTANGLE, MARGIN, top, SLIDE_WIDTH_IN - 2 * MARGIN, LANE_H,
                       fill=LANE_BAND, line=None, name=f"Lane band {lane}")
        fill, border = CATEGORY_STYLE[category]
        faulted = lane == FAULT_LANE
        chip = _add_shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, LANE_LABEL_X + 0.05, top + 0.16, 0.9, LANE_H - 0.32,
                          fill=fill, line=FAILURE if faulted else border, line_pt=1.5 if faulted else 1.0,
                          name=f"Lane {lane}")
        frame = chip.text_frame
        _set_margins(frame, 0.04)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = [(lane, 12, True, FAILURE if faulted else INK)]
        if faulted:
            lines.append(("faulted", 9, False, FAILURE))
        _write_lines(frame, lines, align=PP_ALIGN.CENTER)

    headers = [(FAULT_X, FAULT_W, "FAULT")]
    headers += [(EVENT_X0 + (step - 1) * EVENT_COL_W, EVENT_COL_W, str(step)) for step in range(1, len(RECOVERY_EVENTS) + 1)]
    headers.append((_handoff_x() - HANDOFF_W / 2, HANDOFF_W, "HANDOFF"))
    headers.append((_nominal_x(0) - (NOMINAL_COL_W - NOMINAL_W) / 2, NOMINAL_COL_W * len(NOMINAL_EVENTS),
                    "NOMINAL EVENTS THAT CONNECT TO THE PLAN"))
    for x, w, text in headers:
        _add_text(shapes, x, HEADER_Y, w, 0.24, [(text, 9.5, True, MUTED)], align=PP_ALIGN.CENTER,
                  name=f"Header {text}")


def _draw_diagram(shapes) -> None:
    fault_y = _lane_center(FAULT_LANE)
    fault_cx = FAULT_X + FAULT_W / 2
    fault = _add_shape(shapes, MSO_SHAPE.OVAL, fault_cx - 0.17, fault_y - 0.27, 0.34, 0.34, fill=FAILURE,
                       line=None, name="Fault marker")
    frame = fault.text_frame
    _set_margins(frame, 0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _write_lines(frame, [("×", 16, True, WHITE)], align=PP_ALIGN.CENTER)
    _add_text(shapes, FAULT_X, fault_y + 0.1, FAULT_W, 0.32, [("gear_small", 8.5, True, FAILURE),
                                                              ("dropped", 8.5, False, FAILURE)],
              align=PP_ALIGN.CENTER, name="Fault label")

    for step, event in enumerate(RECOVERY_EVENTS, start=1):
        _add_event_box(shapes, _event_x(step), event, nominal=False, label=f"Recovery event {step}")
    _add_review_flag(shapes, _event_x(FLAGGED_STEP) + EVENT_W - 0.2,
                     _lane_center(RECOVERY_EVENTS[FLAGGED_STEP - 1].lane) - EVENT_H / 2 - 0.14, "Review flag")

    # Sequence lines run at the acting resource's lane, so a box that spans two lanes is entered at its own row.
    _line(shapes, fault_cx + 0.17, fault_y - 0.1, _event_x(1), fault_y - 0.1)
    for step, (earlier, later) in enumerate(zip(RECOVERY_EVENTS, RECOVERY_EVENTS[1:]), start=1):
        earlier_y, later_y = _lane_center(earlier.lane), _lane_center(later.lane)
        if earlier_y == later_y:
            _line(shapes, _event_x(step) + EVENT_W, earlier_y, _event_x(step + 1), later_y)
            continue
        # A lane change is drawn as an L: out of the earlier box toward the later lane, then into the later box.
        drop_x = _event_x(step) + EVENT_W / 2
        edge_y = earlier_y + (EVENT_H / 2 if later_y > earlier_y else -EVENT_H / 2)
        _line(shapes, drop_x, edge_y, drop_x, later_y, arrow=False)
        _line(shapes, drop_x, later_y, _event_x(step + 1), later_y)

    handoff_x = _handoff_x()
    last_y = _lane_center(RECOVERY_EVENTS[-1].lane)
    _line(shapes, _event_x(len(RECOVERY_EVENTS)) + EVENT_W, last_y, handoff_x, last_y, arrow=False)
    _line(shapes, handoff_x, LANE_TOP + 0.04, handoff_x, LANE_TOP + len(LANES) * LANE_H - 0.04, color=INK,
          width_pt=2.0, arrow=False)
    _add_text(shapes, handoff_x - HANDOFF_W / 2, LANE_TOP + len(LANES) * LANE_H + 0.04, HANDOFF_W, 0.24,
              [("connected", 9, False, MUTED)], align=PP_ALIGN.CENTER, name="Handoff status")

    nominal_boxes = [_add_event_box(shapes, _nominal_x(index), event, nominal=True, label=f"Nominal event {index + 1}")
                     for index, event in enumerate(NOMINAL_EVENTS)]
    first_y = _lane_center(NOMINAL_EVENTS[0].lane)
    _line(shapes, handoff_x, first_y, _nominal_x(0), first_y)
    for earlier, later in zip(nominal_boxes, nominal_boxes[1:]):
        _connect(shapes, earlier, RIGHT, later, LEFT)
    end_x = _nominal_x(len(NOMINAL_EVENTS) - 1) + NOMINAL_W
    end_y = _lane_center(NOMINAL_EVENTS[-1].lane)
    _line(shapes, end_x, end_y, end_x + 0.2, end_y)
    _add_text(shapes, end_x + 0.24, end_y - 0.2, SLIDE_WIDTH_IN - MARGIN - end_x - 0.24, 0.4,
              [("remaining", 9.5, False, MUTED), ("plan", 9.5, False, MUTED)], anchor=MSO_ANCHOR.MIDDLE,
              name="Remaining plan")


def _add_card(shapes, x: float, y: float, w: float, h: float, name: str):
    card = _add_shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=CARD_FILL, line=PANEL_EDGE,
                      line_pt=1.0, name=name)
    card.adjustments[0] = 0.05
    return card


def _draw_cards(shapes) -> None:
    top, height, gap = 5.02, 1.98, 0.24
    width = (SLIDE_WIDTH_IN - 2 * MARGIN - 2 * gap) / 3
    xs = [MARGIN + index * (width + gap) for index in range(3)]
    pad = 0.18

    _add_card(shapes, xs[0], top, width, height, "Card failure")
    image_w = 1.9
    image_h = image_w * 312 / 447
    shapes.add_picture(str(SNAPSHOT_PATH), Inches(xs[0] + pad), Inches(top + pad + 0.36), Inches(image_w),
                       Inches(image_h)).name = "Part slippage snapshot"
    _add_text(shapes, xs[0] + pad, top + pad, width - 2 * pad, 0.3, [("Failure", 14, True, INK)], name="Card failure title")
    text_x = xs[0] + pad + image_w + 0.16
    _add_text(shapes, text_x, top + pad + 0.36, xs[0] + width - pad - text_x, height - 0.6,
              [("ur5e-4 drops gear_small during place_approach", 13, False, INK)], name="Card failure text")

    _add_card(shapes, xs[1], top, width, height, "Card created")
    _add_text(shapes, xs[1] + pad, top + pad, width - 2 * pad, height - 2 * pad, [
        ("What the LLM created", 14, True, INK),
        ("4 new events, none in the nominal models", 13, False, INK),
        ("ur5e-3 parks KET4 in the Buffer, picks up gear_small, and assembles it", 13, False, INK),
        ("Then nominal work resumes: ur5e-3 picks KET4 again", 13, False, MUTED),
    ], space_after_pt=5, name="Card created text")

    _add_card(shapes, xs[2], top, width, height, "Card review")
    _add_review_flag(shapes, xs[2] + pad, top + pad + 0.02, "Card review flag")
    _add_text(shapes, xs[2] + pad + 0.38, top + pad, width - 2 * pad - 0.38, 0.3, [("Review findings", 14, True, INK)],
              name="Card review title")
    _add_text(shapes, xs[2] + pad, top + pad + 0.4, width - 2 * pad, height - 2 * pad - 0.4, [
        ("Step 2 puts KET4 in zone_1, but the Buffer's own pick guard needs zone_4_part", 13, False, INK),
        ("The handoff check passed only on ur5e-3's reduced pick model", 13, False, INK),
    ], space_after_pt=5, name="Card review text")


def build() -> Path:
    """Write the slide and return its path."""
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = slide.shapes
    _draw_header(shapes)
    _draw_lanes(shapes)
    _draw_diagram(shapes)
    _draw_cards(shapes)
    slide.notes_slide.notes_text_frame.text = SPEAKER_NOTES
    presentation.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    logger.info("Wrote %s", build())
